from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
OUT = 'test/Test - Profil Excavator.docx'
INK = '1F2430'
RED = 'E53935'
TAN = 'EDCD8B'
SLATE = '4A5162'
GREEN = '2E7D32'
LIGHT = 'F3F4F5'
WHITE = 'FFFFFF'

def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    tc_pr.append(shd)

def set_cell_text(cell, text, bold=False, color=INK, size=9):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = 'Aptos'
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, True, WHITE, 9)
        shade(t.rows[0].cells[i], SLATE)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, False, INK, 8.5)
            if ri % 2 == 1:
                shade(cells[i], LIGHT)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p

def test_case(doc, id_, title, pre, steps, expected):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f'{id_} — {title}')
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(SLATE)
    r.font.size = Pt(10)
    for label, text in [('Préconditions', pre), ('Étapes', steps), ('Résultat attendu', expected)]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.16)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(label + ' : ')
        r.bold = True
        r.font.size = Pt(9)
        r2 = p.add_run(text)
        r2.font.size = Pt(9)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.62)
sec.bottom_margin = Inches(0.62)
sec.left_margin = Inches(0.68)
sec.right_margin = Inches(0.68)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(9)
styles['Normal'].font.color.rgb = RGBColor.from_string(INK)
for name, size, color in [('Title', 24, RED), ('Heading 1', 15, RED), ('Heading 2', 11, SLATE)]:
    st = styles[name]
    st.font.name = 'Aptos Display'
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)

# En-tête et pied de page.
header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
rr = header.add_run('ORE TRACKING  |  PLAN DE TEST')
rr.bold = True
rr.font.size = Pt(8)
rr.font.color.rgb = RGBColor.from_string(SLATE)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
rr = footer.add_run('Profil EXCAVATOR — TAB 3  ·  Document de validation terrain')
rr.font.size = Pt(8)
rr.font.color.rgb = RGBColor.from_string(SLATE)

p = doc.add_paragraph(style='Title')
p.paragraph_format.space_after = Pt(2)
p.add_run('Plan de test — Profil Excavator')
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
r = p.add_run('Version 1.0  |  À remplir pendant les essais sur tablette Android')
r.italic = True
r.font.color.rgb = RGBColor.from_string(SLATE)
r.font.size = Pt(10)

table(doc, ['Campagne', 'Date / heure', 'Opérateur', 'Tablette / ID', 'Version APK'], [['', '', '']], [1.25, 1.25, 1.35, 1.55, 1.35])

heading(doc, '1. Objectif et périmètre', 1)
doc.add_paragraph('Valider le profil EXCAVATOR (TAB 3) sur tablette Android : affichage des mesures GNSS et de la moyenne de la perle, mise en page sans défilement, état BLE/GPS/4G, cycle automatique de session perle à l’entrée/sortie du DT, et bouton « Envoi manuel ».')
doc.add_paragraph('Hors périmètre : alerte hors-parcours sur Excavator, validation de l’affectation, commandes manuelles de début/fin de cycle, et validation de l’écran DT. Le cycle perle doit rester automatique, sans intervention opérateur.')

heading(doc, '2. Matériel et préparation', 1)
table(doc, ['Élément', 'À préparer / noter'], [
    ['Tablette TAB 3', 'Android, profil EXCAVATOR sélectionné, identifiant : __________________'],
    ['Tablette TAB 2', 'DT autorisé dans la base locale, BLE activé, identifiant : __________________'],
    ['GNSS', 'Récepteur/source GNSS disponible ; permissions accordées ; noter le type : __________________'],
    ['Environnement', 'Version APK : __________  ·  base/config : __________  ·  réseau serveur : en ligne / hors ligne'],
    ['Sécurité', 'Essai en zone sûre. Ne pas déplacer la machine ni le camion pour le test sans accord du responsable terrain.'],
])

heading(doc, '3. Règles de verdict', 1)
table(doc, ['Statut', 'Signification'], [
    ['PASS', 'Résultat observé conforme à tous les critères du test.'],
    ['FAIL', 'Résultat non conforme ; consigner l’observation et joindre une capture/photo si possible.'],
    ['BLOQUÉ', 'Test impossible (matériel, permissions, configuration, réseau ou précondition manquante).'],
    ['N/A', 'Test non applicable à cette configuration ; préciser pourquoi.'],
])

doc.add_page_break()
heading(doc, '4. Tests fonctionnels', 1)

test_case(doc, 'EXC-01', 'Ouverture de l’écran et badges', 'La tablette est démarrée sur le profil EXCAVATOR.', 'Ouvrir l’écran principal ; relever le nom affiché et les états 4G, BLE et GPS au repos.', 'L’écran s’ouvre sans erreur ; nom appareil (ou titre de profil de repli) visible ; chaque badge correspond à l’état réel : rouge si arrêté/indisponible, orange en recherche, vert après lien/mesure valide.')
test_case(doc, 'EXC-02', 'Première mesure GNSS et compteurs E/N/Z', 'Source GNSS disponible et autorisations accordées.', 'Attendre une mesure GNSS valide ; comparer E/N/Z à la source de diagnostic disponible ; attendre une mesure suivante.', 'Les tirets initiaux sont remplacés par les coordonnées mesurées ; les valeurs se rafraîchissent à l’arrivée des mesures ; aucune valeur simulée ou fixe n’est présentée comme une mesure réelle.')
test_case(doc, 'EXC-03', 'Aucune mesure GNSS / GPS indisponible', 'GNSS désactivé, permission refusée ou source non disponible (tester le cas contrôlé).', 'Ouvrir l’écran et observer les valeurs et le badge GPS.', 'L’écran reste utilisable ; valeurs absentes indiquées honnêtement (tirets/état défini par l’interface) ; badge GPS non vert sans mesure réelle ; aucune exception non gérée.')
test_case(doc, 'EXC-04', 'Entrée automatique du DT dans le rayon', 'TAB 3 advertising actif ; TAB 2 autorisé, BLE actif ; les appareils sont hors de portée au départ.', 'Approcher TAB 2 de TAB 3 jusqu’à l’authentification ; ne toucher à aucune commande de cycle.', 'TAB 3 indique le DT connecté ; une session perle s’ouvre automatiquement ; compteur de positions démarre à zéro ; aucune action manuelle de cycle n’est demandée.')
test_case(doc, 'EXC-05', 'Accumulation des positions de perle et moyenne', 'Session perle ouverte (EXC-04), GNSS actif.', 'Laisser arriver plusieurs mesures ; relever le compteur de positions et Xmoy/Ymoy/Zmoy.', 'Le nombre de positions augmente avec les mesures reçues pendant la session ; la moyenne s’affiche et se met à jour ; les coordonnées correspondent aux mesures accumulées selon la règle d’arrondi configurée.')
test_case(doc, 'EXC-06', 'Sortie du DT et fermeture automatique', 'Session perle active avec au moins une position accumulée.', 'Éloigner TAB 2 hors du rayon contrôlé, sans toucher à un bouton de cycle ; consulter l’écran et le journal local.', 'La session se ferme automatiquement ; le DT n’est plus indiqué comme présent ; la moyenne de la session est gelée et l’événement MOYENNE_PERLE est enregistré une fois ; la prochaine connexion repart avec un nouvel accumulateur.')
test_case(doc, 'EXC-07', 'Persistance de l’événement et reprise', 'Événements présents en file locale ; lien vers TAB 2 momentanément indisponible.', 'Provoquer/observer une interruption contrôlée, puis rétablir le lien ; comparer les event_id et les événements reçus.', 'Les événements restent persistés tant qu’ils ne sont pas acquittés ; après retour du lien ils peuvent être retransmis ; aucun événement n’est perdu ni dupliqué côté destinataire (idempotence event_id).')
test_case(doc, 'EXC-08', 'Bouton « Envoi manuel » — visibilité et état', 'Écran Excavator ouvert.', 'Observer le bouton sans DT connecté, puis avec DT connecté et file vide, puis avec lien authentifié et éléments non acquittés si ce scénario est disponible.', 'Le bouton est visible dans les états prévus par l’interface ; il est activé uniquement lorsque l’envoi est possible et qu’il reste des données à transmettre ; pendant un envoi, il ne déclenche pas d’envois concurrents.')
test_case(doc, 'EXC-09', 'Bouton « Envoi manuel » — action', 'Un lien BLE TAB 3 ↔ TAB 2 authentifié et des événements en attente sont nécessaires.', 'Appuyer une fois sur « Envoi manuel » ; observer l’indicateur d’envoi, les ACK et l’état de la file.', 'La retransmission est déclenchée ; l’indicateur/flèche reflète l’envoi et l’attente des ACK ; les événements acquittés quittent la file ; un retour visible/journalisé confirme l’action. Aucun effet indésirable si l’appui est répété pendant l’envoi.')
test_case(doc, 'EXC-10', 'Envoi manuel sans lien', 'Lien BLE non authentifié ; données conservées localement si possible.', 'Tenter l’action si le bouton est activable, ou vérifier qu’il est désactivé ; consulter le journal.', 'Aucun événement n’est déclaré transmis ; les données restent en file ; l’interface explique clairement l’absence de lien ou empêche l’action.')
test_case(doc, 'EXC-11', 'État 4G et synchronisation serveur', 'Tester une fois en réseau disponible et une fois sans réseau si possible.', 'Comparer le badge 4G au résultat réel de la synchronisation/diagnostic serveur ; rétablir le réseau.', 'Le badge distingue la disponibilité réseau de la réponse réelle du serveur selon la logique définie ; le retour du réseau ne bloque pas l’affichage GNSS/BLE local.')
test_case(doc, 'EXC-12', 'Pas d’alerte hors-parcours sur Excavator', 'Écran Excavator actif ; aucune exigence d’alerte hors-parcours sur ce profil.', 'Observer l’écran pendant les essais, sans déclencher volontairement d’alerte dangereuse.', 'Aucun bandeau ou signal hors-parcours n’est exigé sur l’écran Excavator. Le test ne pénalise pas l’absence de cette alerte : elle relève de l’écran DT.')

heading(doc, '5. Tests d’affichage — tablette, sans scroll', 1)
doc.add_paragraph('Exécuter sur l’orientation réellement utilisée en exploitation (priorité paysage). Pour chaque taille, vérifier que tout le contenu reste dans la zone entre l’en-tête commun et le bandeau inférieur. Aucun défilement vertical ou horizontal n’est autorisé.')
table(doc, ['ID', 'Résolution logique / orientation', 'Critères de réussite', 'Verdict / notes'], [
    ['LAY-01', 'Tablette terrain — paysage : ______ × ______', 'En-tête, badges, tableau E/N/Z, tableau de moyenne et bouton « Envoi manuel » entièrement visibles ; aucun scroll ; aucun overflow ; textes lisibles.', ''],
    ['LAY-02', 'Tablette terrain — portrait : ______ × ______', 'Même contenu visible sans défilement ; bouton accessible ; tableaux non superposés ; textes essentiels lisibles.', ''],
    ['LAY-03', 'Rotation paysage ↔ portrait', 'Après rotation, layout recalculé ; aucune zone blanche anormale, coupure ou scroll ; données restent affichées.', ''],
    ['LAY-04', 'Valeurs GNSS longues / identifiants longs', 'E/N/Z et DT ID n’empiètent pas sur les badges ou les colonnes ; réduction/ellipse maîtrisée ; aucun overflow.', ''],
])

heading(doc, '6. Résumé d’exécution', 1)
table(doc, ['Résultat', 'Nombre'], [['PASS', ''], ['FAIL', ''], ['BLOQUÉ', ''], ['N/A', '']])

heading(doc, '7. Anomalies et observations', 1)
table(doc, ['N°', 'Test', 'Description / étapes pour reproduire', 'Gravité', 'Capture / référence'], [['', '', ''], ['', '', ''], ['', '', '']])

heading(doc, '8. Décision de validation', 1)
table(doc, ['Décision', 'Cocher / compléter'], [
    ['☐ Accepté pour les essais terrain élargis', 'Aucun défaut bloquant ; tous les scénarios critiques PASS.'],
    ['☐ Accepté avec réserves', 'Réserves / limites : ______________________________________________'],
    ['☐ Refusé — corrections nécessaires', 'Défauts bloquants : ______________________________________________'],
])
doc.add_paragraph('Nom et signature du testeur : ______________________________    Date : ______________')
doc.add_paragraph('Nom et signature du responsable : __________________________    Date : ______________')

# Propriétés du document.
doc.core_properties.title = 'Test - Profil Excavator'
doc.core_properties.subject = 'Plan de test fonctionnel et terrain du profil EXCAVATOR TAB 3'
doc.core_properties.author = 'ORE Tracking'
doc.core_properties.keywords = 'Excavator, TAB 3, test, GNSS, BLE, envoi manuel, responsive'
doc.save(OUT)
print(f'Document créé : {OUT}')
