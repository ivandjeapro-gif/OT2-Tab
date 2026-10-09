# -*- coding: utf-8 -*-
"""Genere la presentation ORE Tracking (5 min, 6 diapos).

Theme inspire de l'app : fond blanc, encadre rouge, tan/or accent.
Sortie : E:\\...\\presentation_ore_tracking.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# Couleurs du projet (AppColors)
ROUGE = RGBColor(0xE5, 0x39, 0x35)
INK = RGBColor(0x1F, 0x24, 0x30)
TAN = RGBColor(0xED, 0xCD, 0x8B)
GRIS = RGBColor(0x4A, 0x51, 0x62)
BLANC = RGBColor(0xFF, 0xFF, 0xFF)
VERT = RGBColor(0x72, 0xC8, 0x77)
ORANGE = RGBColor(0xE8, 0x87, 0x1E)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BLANK = prs.slide_layouts[6]


def ajout_fond(slide, couleur=BLANC):
    fond = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    fond.fill.solid()
    fond.fill.fore_color.rgb = couleur
    fond.line.fill.background()
    fond.shadow.inherit = False
    slide.shapes._spTree.remove(fond._element)
    slide.shapes._spTree.insert(2, fond._element)


def bandeau(slide, couleur=ROUGE, hauteur=Inches(0.35)):
    bande = slide.shapes.add_shape(1, 0, 0, prs.slide_width, hauteur)
    bande.fill.solid()
    bande.fill.fore_color.rgb = couleur
    bande.line.fill.background()
    bande.shadow.inherit = False


def titre(slide, texte, taille=34):
    boite = slide.shapes.add_textbox(Inches(0.6), Inches(0.55), Inches(12.1), Inches(1.0))
    tf = boite.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = texte
    r.font.size = Pt(taille)
    r.font.bold = True
    r.font.color.rgb = INK
    return boite


def filet(slide):
    ligne = slide.shapes.add_shape(1, Inches(0.65), Inches(1.55), Inches(3.2), Pt(3))
    ligne.fill.solid()
    ligne.fill.fore_color.rgb = TAN
    ligne.line.fill.background()
    ligne.shadow.inherit = False


def puces(slide, items, gauche=0.7, haut=1.85, largeur=12.0, taille=20, interligne=8):
    boite = slide.shapes.add_textbox(Inches(gauche), Inches(haut), Inches(largeur), Inches(5.0))
    tf = boite.text_frame
    tf.word_wrap = True
    for i, (texte, niveau) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(interligne)
        if isinstance(texte, str):
            texte = [(texte, False)]
        for morceau, gras in texte:
            r = p.add_run()
            r.text = ("•  " if niveau == 0 else "–  ") + morceau if niveau >= 0 else morceau
            r.font.size = Pt(taille if niveau == 0 else taille - 3)
            r.font.bold = gras
            r.font.color.rgb = INK if niveau >= 0 else GRIS
    return boite


def pied(slide, numero):
    b = slide.shapes.add_textbox(Inches(12.3), Inches(7.0), Inches(0.8), Inches(0.4))
    p = b.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = f"{numero} / 6"
    r.font.size = Pt(12)
    r.font.color.rgb = GRIS


# ---------------------------------------------------------------- Diapo 1
s = prs.slides.add_slide(BLANK)
ajout_fond(s)
bandeau(s)
titre(s, "ORE Tracking — Suivre le minerai, même hors ligne")
filet(s)
puces(s, [
    ("Le besoin terrain : tracer le minerai du blast jusqu'à la dumping area", True),
    ("Site minier : réseau cellulaire instable, papier et radios — aucune traçabilité fiable", False),
    ("Conséquences : erreurs de destination, re-saisies, litiges de comptabilité minière", False),
    ("ORE Tracking = 3 tablettes durcies + un serveur, conçues pour travailler SANS réseau", False),
    ("Principe directeur : « rien ne se perd, tout se rejoue »", True),
], taille=22, interligne=12)
pied(s, 1)

# ---------------------------------------------------------------- Diapo 2
s = prs.slides.add_slide(BLANK)
ajout_fond(s)
bandeau(s, couleur=GRIS)
titre(s, "La solution : 3 tablettes, 1 serveur")
filet(s)
# Trois encarts
for i, (t, sous_t, couleur) in enumerate([
    ("TAB 1 — Spotter", "Saisit le blast : blocs, lignes, photos", ORANGE),
    ("TAB 2 — DT (camion)", "Passerelle : reçoit tout, relaie au serveur", VERT),
    ("TAB 3 — Excavatrice", "Position GPS du point de chargement", ROUGE),
]):
    x = Inches(0.7 + i * 4.1)
    carte = s.shapes.add_shape(1, x, Inches(1.9), Inches(3.8), Inches(2.6))
    carte.fill.solid()
    carte.fill.fore_color.rgb = BLANC
    carte.line.color.rgb = couleur
    carte.line.width = Pt(2.5)
    carte.shadow.inherit = False
    tf = carte.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = t
    r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = couleur
    p2 = tf.add_paragraph(); p2.space_before = Pt(8)
    r2 = p2.add_run(); r2.text = sous_t
    r2.font.size = Pt(15); r2.font.color.rgb = INK
puces(s, [
    ("Le DT est le pivot : point d'accès Wi-Fi pour le spotter, passerelle BLE pour l'excavatrice", True),
    ("Chaque tablette garde sa base SQLite locale : l'application fonctionne intégralement hors ligne", True),
], haut=4.9, taille=20)
pied(s, 2)

# ---------------------------------------------------------------- Diapo 3
s = prs.slides.add_slide(BLANK)
ajout_fond(s)
bandeau(s)
titre(s, "Trois liens, un seul objectif : ne jamais perdre une donnée")
filet(s)
puces(s, [
    ("BLE (Bluetooth) — excavatrice → DT : positions GPS continues, faible consommation", True),
    ("Wi-Fi hotspot — spotter → DT : le blast complet (lignes + photos) transmis au verrouillage", True),
    ("4G — DT → serveur : synchronisation en tâche de fond dès que le réseau revient", True),
    ("File PENDING persistante : chaque événement attend son accusé de réception", False),
    ("Reconnexions automatiques + retransmission de la file : une coupure n'est jamais une perte", False),
    ("Idempotence : rejouer un envoi ne duplique rien côté serveur", False),
], taille=21, interligne=10)
pied(s, 3)

# ---------------------------------------------------------------- Diapo 4
s = prs.slides.add_slide(BLANK)
ajout_fond(s)
bandeau(s, couleur=GRIS)
titre(s, "L'application aujourd'hui (captures)")
filet(s)
puces(s, [
    ("Écran Spotter : saisie du blast, photo du bloc, verrouillage puis envoi", True),
    ("Écran DT : carte de suivi, destination reçue du serveur, alerte hors-parcours (40 m)", True),
    ("Voyants temps réel : 4G — BLE — Wi-Fi — GPS, chacun rouge / orange / vert", True),
    ("Barre du bas : flèche animée = envoi vers le DT en cours ; nuage animé = synchro serveur", True),
    ("[INSÉRER ICI : capture écran DT]", False),
    ("[INSÉRER ICI : capture écran spotter]", False),
], taille=20, interligne=8)
pied(s, 4)

# ---------------------------------------------------------------- Diapo 5
s = prs.slides.add_slide(BLANK)
ajout_fond(s)
bandeau(s)
titre(s, "Fiabilité : testée avant d'être déployée")
filet(s)
puces(s, [
    ("Architecture injectable : le même code tourne sur un banc simulé ou sur les vraies piles réseau", True),
    ("82 tests verts : protocole, fragmentation TCP, coupures, migrations base, file PENDING", True),
    ("Transports réels validés en boucle locale : handshake + ACK bout en bout, JSON fragmenté réassemblé", True),
    ("SQLite migrations idempotentes : un crash pendant une mise à jour ne corrompt rien", True),
    ("Web & Windows : la même base tourne sur tablette, navigateur et PC (développement)", False),
], taille=21, interligne=10)
pied(s, 5)

# ---------------------------------------------------------------- Diapo 6
s = prs.slides.add_slide(BLANK)
ajout_fond(s)
bandeau(s, couleur=GRIS)
titre(s, "Prochaines étapes")
filet(s)
puces(s, [
    ("Validation sur tablettes réelles : BLE + Wi-Fi entre deux appareils Android", True),
    ("Barre de trajet branchée sur le GPS réel (aujourd'hui simulée côté affichage)", True),
    ("GPS externe USB-C : voyant en place, récepteur à tester sur le camion", True),
    ("Détection d'écart déjà réelle : alerte si le DT sort de 40 m du parcours pendant 15 s", False),
    ("Merci — questions ?  |  ORE Tracking · Lafigue", True),
], taille=21, interligne=10)
pied(s, 6)

sortie = r"E:\PROJETS\VISION MINING\DB\OT V2\OT 2 - Tab\presentation_ore_tracking.pptx"
prs.save(sortie)
print("OK :", sortie)
